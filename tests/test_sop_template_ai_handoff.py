from __future__ import annotations

import json
import base64
import tempfile
import unittest
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE

from scripts.generate_sop_template_ai_handoff import (
    CENTER_FLOWCHART_NAME,
    CONTENT_PROFILE_HDMI,
    FINAL_DOCX_NAME,
    FORMAT_CHECK_NAME,
    HDMI_FINAL_DOCX_NAME,
    HDMI_TEMPLATE_ID,
    MANIFEST_NAME,
    TEMPLATE_ID,
    VALIDATION_NAME,
    _group_methods_for_slots,
    generate_package,
    generate_route_package,
    validate_document,
)
from cad_ai.sop_visual_template import _work_image_body_row_heights
from cad_ai.sop_knowledge.store import SopKnowledgeStore
from cad_ai.sop_knowledge.models import RouteSectionDraft
from cad_ai.sop_knowledge.documents import SopDocumentService
from cad_ai.sop_visual_template import _work_image_font_sizes
from tests.test_sop_knowledge_workflow import make_identity, make_route


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class SopTemplateAiHandoffTests(unittest.TestCase):
    def test_three_slot_layout_reserves_space_for_complete_signoff_table(self) -> None:
        heights = _work_image_body_row_heights(3, caption_line_count=2, font_profile="standard")
        self.assertLessEqual(sum(heights), 3980)

    def test_image_caption_grouping_preserves_explicit_blank_slots(self) -> None:
        self.assertEqual(
            _group_methods_for_slots(["准备材料", "", "", "接通设备", "", "记录结果"], 6),
            ["准备材料", "", "", "接通设备", "", "记录结果"],
        )

    def test_frozen_handoff_entrypoint_generates_exact_two_section_template(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = generate_package(directory, document_date="2026-08-11")
            root = Path(directory)

            expected_files = {
                FINAL_DOCX_NAME,
                CENTER_FLOWCHART_NAME,
                MANIFEST_NAME,
                FORMAT_CHECK_NAME,
                VALIDATION_NAME,
            }
            self.assertEqual({path.name for path in root.iterdir()}, expected_files)
            self.assertEqual(result["template_id"], TEMPLATE_ID)
            self.assertTrue(result["structural_pass"])
            self.assertTrue(result["visual_qa_required"])

            document = Document(root / FINAL_DOCX_NAME)
            self.assertEqual(len(document.sections), 2)
            self.assertEqual(len(document.tables), 8)
            self.assertEqual(document.tables[0].cell(2, 3).text.strip(), "DRAFT")
            self.assertEqual(document.tables[4].cell(2, 3).text.strip(), "DRAFT")
            self.assertEqual(document.tables[0].cell(2, 5).text.strip(), "2026/8/11")
            self.assertEqual(document.tables[4].cell(1, 7).text.strip(), "2026/8/11")
            self.assertEqual(
                [document.tables[7].cell(1, index).text.strip() for index in range(3)],
                ["", "", ""],
            )

            validation = json.loads((root / VALIDATION_NAME).read_text(encoding="utf-8"))
            self.assertTrue(validation["structural_pass"])
            self.assertEqual(validation["visual_qa"]["expected_page_count"], 2)

            manifest = json.loads((root / MANIFEST_NAME).read_text(encoding="utf-8"))
            self.assertEqual(manifest["template_id"], TEMPLATE_ID)
            self.assertTrue(manifest["tables_filled_before_flowchart"])
            self.assertEqual(manifest["fixed_template_profile"]["step_order"], "1,2,3 / 6,5,4")
            self.assertTrue(manifest["guardrails"]["no_auto_release"])

    def test_hdmi_profile_rejects_the_obsolete_two_page_route(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "route-backed multi-page"):
                generate_package(
                    directory,
                    document_date="2026-08-11",
                    content_profile=CONTENT_PROFILE_HDMI,
                )

    def test_route_backed_hdmi_generates_one_flow_page_and_repeated_instruction_pages(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = SopKnowledgeStore(root / "knowledge.sqlite3")
            store.initialize()
            store.ensure_process_family("test_family", "测试工艺族")
            identity = make_identity("HDMI-ROUTE-TEST")
            store.upsert_product(identity, {"class": "cable"})
            route_id = store.create_route(make_route(identity, 3))
            store.create_route_section(
                route_id,
                RouteSectionDraft(
                    section_type="ie_timing",
                    content={"单价": "15.50", "人数": "2", "source": "人工填写"},
                ),
            )
            step = store.get_route(route_id)["steps"][1]
            store.update_step_field(
                step["id"],
                "record_output",
                ["首件记录", "巡检记录", "登记工单号和异常现象"],
                reviewer="worker-01",
                decision="needs_revision",
            )
            result = generate_route_package(
                root / "package",
                document_date="2026-08-12",
                db_path=store.path,
                route_id=route_id,
            )

            self.assertEqual(result["template_id"], HDMI_TEMPLATE_ID)
            self.assertEqual(result["instruction_page_count"], 3)
            self.assertEqual(result["expected_page_count"], 4)
            document = Document(result["document_docx"])
            self.assertEqual(len(document.sections), 2)
            self.assertEqual(len(document.tables), 16)
            first_instruction_header = document.tables[4]
            first_instruction_body = document.tables[5]
            first_instruction_detail = document.tables[6]
            first_instruction_footer = document.tables[7]
            self.assertAlmostEqual(first_instruction_body.columns[0].width.cm, 1.1, places=1)
            self.assertAlmostEqual(first_instruction_body.columns[14].width.cm, 1.74, places=1)
            self.assertGreaterEqual(first_instruction_header.rows[0].height.twips, 360)
            self.assertGreaterEqual(first_instruction_header.rows[1].height.twips, 360)
            self.assertGreaterEqual(first_instruction_header.rows[2].height.twips, 360)
            self.assertTrue(
                all(row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST for row in first_instruction_body.rows)
            )
            self.assertTrue(
                all(row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST for row in first_instruction_detail.rows)
            )
            self.assertTrue(
                all(row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY for row in first_instruction_footer.rows)
            )
            body_font_sizes = [
                run.font.size.pt
                for row in first_instruction_body.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            ]
            detail_font_sizes = [
                run.font.size.pt
                for row in first_instruction_detail.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            ]
            self.assertGreaterEqual(min(body_font_sizes), 7.5)
            self.assertGreaterEqual(min(detail_font_sizes), 7.5)
            self.assertEqual([step["work_image_slots"] for step in store.get_route(route_id)["steps"]], [3, 3, 3])
            first_body_text = "\n".join(cell.text for row in first_instruction_body.rows for cell in row.cells)
            self.assertIn("作业方法（Operating method）", first_body_text)
            self.assertIn("工程名称（Project name）", first_body_text)
            self.assertIn("使用材料（Use material）", first_body_text)
            self.assertIn("使用工治具（The use of tooling）", first_body_text)
            self.assertIn("工艺参数（Process parameter）", first_body_text)
            flow_ie_time = document.tables[2]
            self.assertIn("单价", flow_ie_time.cell(1, 3).text)
            self.assertIn("人数", flow_ie_time.cell(1, 4).text)
            self.assertIn("15.50", flow_ie_time.cell(2, 3).text)
            self.assertIn("2", flow_ie_time.cell(2, 4).text)
            for page_index in range(3):
                base = 4 + page_index * 4
                self.assertEqual(document.tables[base].cell(1, 17).text.strip(), "DRAFT")
                work_body_text = "\n".join(
                    cell.text for row in document.tables[base + 1].rows for cell in row.cells
                )
                self.assertIn("IE 工时（人工填写）", work_body_text)
                self.assertIn("单价 15.50", work_body_text)
                self.assertIn("人数 2", work_body_text)
                self.assertEqual(
                    [document.tables[base + 3].cell(0, index).text.strip() for index in range(3)],
                    ["核准：", "审核：", "制表："],
                )
            second_instruction_detail = document.tables[10]
            second_detail_text = "\n".join(
                cell.text for row in second_instruction_detail.rows for cell in row.cells
            )
            self.assertIn("注意事项", second_detail_text)
            self.assertIn("登记工单号和异常现象", second_detail_text)

    def test_route_backed_hdmi_allows_a_single_instruction_page(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = SopKnowledgeStore(root / "knowledge.sqlite3")
            store.initialize()
            store.ensure_process_family("test_family", "Test process family")
            identity = make_identity("HDMI-SINGLE-STEP")
            store.upsert_product(identity, {"class": "cable"})
            route_id = store.create_route(make_route(identity, 1))

            result = generate_route_package(
                root / "package",
                document_date="2026-08-12",
                db_path=store.path,
                route_id=route_id,
            )

            self.assertTrue(result["structural_pass"])
            self.assertEqual(result["instruction_page_count"], 1)
            self.assertEqual(result["expected_page_count"], 2)
            document = Document(result["document_docx"])
            self.assertEqual(len(document.sections), 2)
            self.assertEqual(len(document.tables), 8)
            validation = json.loads((root / "package" / VALIDATION_NAME).read_text(encoding="utf-8"))
            self.assertTrue(validation["structural_pass"])
            self.assertEqual(validation["errors"], [])

    def test_route_backed_hdmi_supports_every_work_image_layout(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = SopKnowledgeStore(root / "knowledge.sqlite3")
            store.initialize()
            store.ensure_process_family("test_family", "测试工艺族")
            identity = make_identity("HDMI-LAYOUT-TEST")
            store.upsert_product(identity, {"class": "cable"})
            route_id = store.create_route(make_route(identity, 6))
            steps = store.get_route(route_id)["steps"]
            store.update_step_field(
                steps[0]["id"],
                "method",
                ["准备物料", "核对方向", "执行作业", "记录结果"],
                reviewer="layout-tester",
            )
            for field_name, value in {
                "quality_check": ["本工序检查方法尚未由人工提供。"],
                "acceptance_criteria": ["本工序合格判据尚未由责任人依据受控规范提供。"],
                "tool_equipment": ["设备、工具及治具型号待工程确认。"],
                "safety": ["信息不完整或结果异常时停止流转并提交人工判定。"],
                "record_output": ["本工序记录要求尚未由人工提供。"],
                "inputs": ["本工序输入资料待责任人确认。"],
            }.items():
                store.update_step_field(
                    steps[-1]["id"], field_name, value, reviewer="layout-tester"
                )
            store.update_step_field(
                steps[-1]["id"],
                "method",
                ["准备材料", "", "", "接通设备", "", "记录结果"],
                reviewer="layout-tester",
            )
            for slot_count, step in enumerate(steps, start=1):
                store.set_step_work_image_slots(step["id"], slot_count, reviewer="layout-tester")
            store.set_route_font_profile(route_id, "large", reviewer="layout-tester")

            result = generate_route_package(
                root / "package",
                document_date="2026-08-18",
                db_path=store.path,
                route_id=route_id,
            )

            document = Document(result["document_docx"])
            for page_index, slot_count in enumerate(range(1, 7)):
                body = document.tables[5 + page_index * 4]
                detail_table = document.tables[6 + page_index * 4]
                footer = document.tables[7 + page_index * 4]
                self.assertEqual(footer.rows[0].height_rule, WD_ROW_HEIGHT_RULE.EXACTLY)
                body_text = "\n".join(cell.text for row in body.rows for cell in row.cells)
                for image_number in range(1, slot_count + 1):
                    self.assertIn(f"图 {image_number}", body_text)
                self.assertEqual(len(detail_table.rows), 5)
                self.assertIn("工艺参数", body_text)
                self.assertIn("IE 工时（人工填写）", body_text)
                body_sizes = [
                    run.font.size.pt
                    for row in body.rows
                    for cell in row.cells
                    for paragraph in cell.paragraphs
                    for run in paragraph.runs
                    if run.font.size is not None
                ]
                detail_sizes = [
                    run.font.size.pt
                    for row in detail_table.rows
                    for cell in row.cells
                    for paragraph in cell.paragraphs
                    for run in paragraph.runs
                    if run.font.size is not None
                ]
                self.assertGreaterEqual(min(body_sizes), 8.0)
                self.assertGreaterEqual(min(detail_sizes), 8.0)

            first_body_text = "\n".join(
                cell.text for row in document.tables[5].rows for cell in row.cells
            )
            for method in ("准备物料", "核对方向", "执行作业", "记录结果"):
                self.assertIn(method, first_body_text)

            six_slot_body = document.tables[5 + 5 * 4]
            six_slot_text = "\n".join(cell.text for row in six_slot_body.rows for cell in row.cells)
            self.assertIn("1. 准备材料", six_slot_text)
            self.assertIn("图 2\n未配图", six_slot_text)
            self.assertIn("4. 接通设备", six_slot_text)
            self.assertIn("6. 记录结果", six_slot_text)

            manifest = json.loads((root / "package" / MANIFEST_NAME).read_text(encoding="utf-8"))
            self.assertEqual(
                [item["work_image_slots"] for item in manifest["layout"]["instruction_layouts"]],
                [1, 2, 3, 4, 5, 6],
            )
            validation = json.loads((root / "package" / VALIDATION_NAME).read_text(encoding="utf-8"))
            self.assertTrue(validation["checks"]["manual_ie_area_every_page"])
            self.assertTrue(validation["checks"]["operation_sheet_sections_every_page"])
            self.assertTrue(validation["checks"]["visual_step_order_every_page"])

            documents = SopDocumentService(store)
            if documents._find_libreoffice_executable() is not None:
                preview_dir = root / "libreoffice-preview"
                pdf_path = documents._convert_docx_with_libreoffice(
                    Path(result["document_docx"]), preview_dir
                )
                page_paths = documents._render_pdf_pages(pdf_path, preview_dir)
                with pymupdf.open(pdf_path) as pdf:
                    page_summaries = [
                        " | ".join(page.get_text().splitlines()[:8])
                        for page in pdf
                    ]
                self.assertEqual(
                    len(page_paths),
                    result["expected_page_count"],
                    page_summaries,
                )

    def test_route_backed_hdmi_writes_manual_step_ie_items_without_inventing_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = SopKnowledgeStore(root / "knowledge.sqlite3")
            store.initialize()
            store.ensure_process_family("test_family", "测试工艺族")
            identity = make_identity("HDMI-IE-ITEM-TEST")
            store.upsert_product(identity, {"class": "cable"})
            route_id = store.create_route(make_route(identity, 2))
            step = store.get_route(route_id)["steps"][0]
            store.replace_step_ie_items(
                step["id"],
                [
                    {
                        "action": "剥皮",
                        "machine_type": "人工填写的剥皮机",
                        "equipment_speed": "每分钟 12 米",
                        "headcount": "1",
                        "standard_time": "30 秒",
                        "allowance_rate": "8%",
                        "standard_capacity": "100 条/小时",
                        "time_source": "现场实测记录 IE-2026-08",
                        "note": "单价尚未提供",
                    },
                    {"action": "检查"},
                ],
                reviewer="ie-reviewer",
            )

            result = generate_route_package(
                root / "package",
                document_date="2026-08-21",
                db_path=store.path,
                route_id=route_id,
            )
            document = Document(result["document_docx"])
            body_text = "\n".join(
                cell.text for row in document.tables[5].rows for cell in row.cells
            )
            self.assertIn("IE 工时（人工填写）", body_text)
            self.assertIn("1. 剥皮", body_text)
            self.assertIn("机器类型 人工填写的剥皮机", body_text)
            self.assertIn("设备速度 每分钟 12 米", body_text)
            self.assertNotIn("单价 待确认", body_text)
            self.assertIn("标准产能 100 条/小时", body_text)
            self.assertIn("2. 检查", body_text)

    def test_route_backed_hdmi_embeds_only_confirmed_step_media(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = SopKnowledgeStore(root / "knowledge.sqlite3")
            store.initialize()
            store.ensure_process_family("test_family", "测试工艺族")
            identity = make_identity("HDMI-MEDIA-TEST")
            store.upsert_product(identity, {"class": "cable"})
            route_id = store.create_route(make_route(identity, 2))
            steps = store.get_route(route_id)["steps"]
            confirmed = store.upload_media_asset(
                route_id, original_name="confirmed.png", mime_type="image/png",
                data=PNG_1X1, uploaded_by="worker-01",
            )
            draft = store.upload_media_asset(
                route_id, original_name="draft.png", mime_type="image/png",
                data=PNG_1X1 + b"draft", uploaded_by="worker-01",
            )
            store.link_media_asset(steps[0]["id"], confirmed["id"], caption="确认图片")
            store.confirm_step(steps[0]["id"], reviewer="worker-01")
            store.link_media_asset(steps[1]["id"], draft["id"], caption="草稿图片")

            result = generate_route_package(
                root / "package", document_date="2026-08-12", db_path=store.path, route_id=route_id,
            )

            with zipfile.ZipFile(result["document_docx"]) as archive:
                embedded = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(embedded), 2)  # 流程图 + 1 张已确认工序图片
            document = Document(result["document_docx"])
            first_page_body = document.tables[5]
            second_page_body = document.tables[9]
            first_page_text = "\n".join(cell.text for row in first_page_body.rows for cell in row.cells)
            second_page_text = "\n".join(cell.text for row in second_page_body.rows for cell in row.cells)
            self.assertLess(first_page_text.count("未配图"), second_page_text.count("未配图"))

    def test_check_only_validation_rejects_missing_document(self) -> None:
        result = validate_document(Path("does-not-exist.docx"))
        self.assertFalse(result["structural_pass"])
        self.assertEqual(result["errors"], ["document_not_found"])


if __name__ == "__main__":
    unittest.main()
