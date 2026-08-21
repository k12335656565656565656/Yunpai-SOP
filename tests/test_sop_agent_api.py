from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pydantic import ValidationError

from cad_ai.sop_agent import SOP_GENERATION_SEQUENCE, SopGenerateRequest, SopRoutingStep, generate_sop_package
from cad_ai.sop_agent import _build_model_prompt, _build_structured_sop_data, _ollama_native_chat_payload, _ollama_native_chat_url
from cad_ai.sop_api import create_sop_fastapi_app


class SopAgentApiTests(unittest.TestCase):
    def test_sop_parameter_schema_rejects_nonpositive_production_values(self) -> None:
        with self.assertRaises(ValidationError):
            SopGenerateRequest(product_name="测试产品", part_no="TEST-001", document_no="SOP-001", speed_m_per_min=0)
        with self.assertRaises(ValidationError):
            SopRoutingStep(name="成型", cavity_count=-1)

    def test_human_parameters_are_grounded_in_json_prompt_and_word(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            request = _sample_request(Path(directory)).model_copy(
                update={
                    "speed_m_per_min": 18.5,
                    "mold_id": "MOLD-USB-08",
                    "cavity_count": 4,
                    "tech_params": {"规格": "USB-C 1.0 m", "公差": "±0.5 mm", "材质": "PVC"},
                }
            )
            request.routing_steps[0] = request.routing_steps[0].model_copy(
                update={"speed_m_per_min": 12.0, "tech_params": {"材质": "阻燃 PVC"}}
            )

            response = generate_sop_package(request)
            parsed = json.loads(response.artifacts.parsed_sop_json.read_text(encoding="utf-8"))

            self.assertEqual(parsed["production_parameters"]["speed_m_per_min"], 18.5)
            self.assertEqual(parsed["production_parameters"]["mold_id"], "MOLD-USB-08")
            self.assertEqual(parsed["production_parameters"]["cavity_count"], 4)
            self.assertEqual(parsed["technical_parameters"]["values"]["材质"], "PVC")
            self.assertEqual(parsed["step_parameters"][0]["production_parameters"]["speed_m_per_min"], 12.0)
            self.assertEqual(parsed["step_parameters"][0]["production_parameters"]["mold_id"], "MOLD-USB-08")
            self.assertEqual(parsed["step_parameters"][0]["technical_parameters"]["values"]["材质"], "阻燃 PVC")

            prompt = _build_model_prompt(request)
            self.assertIn("18.5 米/分钟", prompt)
            self.assertIn("MOLD-USB-08", prompt)
            self.assertIn("不得推测、补写或改写", prompt)

            document_text = _document_table_text(Document(response.artifacts.document_docx))
            for expected in ["技术参数", "生产参数", "USB-C 1.0 m", "±0.5 mm", "PVC", "18.5 米/分钟", "MOLD-USB-08", "4"]:
                self.assertIn(expected, document_text)

    def test_missing_parameters_remain_pending_in_json_and_word(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            response = generate_sop_package(_sample_request(Path(directory)))
            parsed = json.loads(response.artifacts.parsed_sop_json.read_text(encoding="utf-8"))

            self.assertEqual(parsed["technical_parameters"]["status"], "needs_confirmation")
            self.assertEqual(parsed["production_parameters"]["status"], "needs_confirmation")
            self.assertIsNone(parsed["production_parameters"]["speed_m_per_min"])
            self.assertIsNone(parsed["production_parameters"]["mold_id"])
            self.assertIsNone(parsed["production_parameters"]["cavity_count"])

            document_text = _document_table_text(Document(response.artifacts.document_docx))
            self.assertIn("技术参数", document_text)
            self.assertIn("生产参数", document_text)
            self.assertGreaterEqual(document_text.count("待确认"), 3)

    def test_model_cannot_inject_ungrounded_sop_parameters(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            request = _sample_request(Path(directory)).model_copy(update={"use_model": True})
            fabricated = _build_structured_sop_data(request) | {
                "technical_parameters": {"values": {"材质": "AI虚构材料"}},
                "production_parameters": {
                    "speed_m_per_min": 999,
                    "mold_id": "AI-MOLD",
                    "cavity_count": 99,
                },
            }

            with patch("cad_ai.sop_agent._generate_content_with_model", return_value=(fabricated, {})):
                response = generate_sop_package(request)

            parsed = json.loads(response.artifacts.parsed_sop_json.read_text(encoding="utf-8"))
            self.assertNotIn("AI虚构材料", json.dumps(parsed, ensure_ascii=False))
            self.assertIsNone(parsed["production_parameters"]["speed_m_per_min"])
            self.assertIsNone(parsed["production_parameters"]["mold_id"])
            self.assertIsNone(parsed["production_parameters"]["cavity_count"])

    def test_ollama_native_payload_disables_thinking_for_qwen35b_json(self) -> None:
        url = _ollama_native_chat_url("http://127.0.0.1:11434/v1/chat/completions")
        payload = _ollama_native_chat_payload(
            {
                "model": "qwen3.6:35b",
                "messages": [{"role": "user", "content": "Return JSON."}],
                "max_tokens": 1900,
                "temperature": 0.1,
            }
        )

        self.assertEqual(url, "http://127.0.0.1:11434/api/chat")
        self.assertEqual(payload["model"], "qwen3.6:35b")
        self.assertFalse(payload["think"])
        self.assertEqual(payload["format"], "json")
        self.assertEqual(payload["options"]["num_predict"], 1900)

    def test_generate_sop_package_from_structured_request(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            request = _sample_request(Path(directory))

            response = generate_sop_package(request)

            self.assertEqual(response.status, "demo_not_for_release")
            self.assertEqual(response.generation_sequence, SOP_GENERATION_SEQUENCE)
            self.assertTrue(response.tables_filled_before_flowchart)
            self.assertEqual(response.center_flowchart_target, "process_flow_body_table_cell_0_0")
            self.assertTrue(response.artifacts.document_docx.exists())
            self.assertTrue(response.artifacts.center_flowchart_png.exists())
            self.assertTrue(response.artifacts.manifest_json.exists())
            self.assertEqual(response.validation["sections"], 2)
            self.assertEqual(response.validation["top_level_tables"], 8)
            self.assertTrue(response.validation["has_png_media"])
            self.assertFalse(response.validation["has_svg"])
            self.assertFalse(response.validation["has_vml_shape"])
            self.assertFalse(response.validation["contains_replacement_char"])
            self.assertTrue(response.validation["contains_ie_time_title"])
            self.assertTrue(response.validation["contains_machine_model_field"])

            document = Document(response.artifacts.document_docx)
            self.assertEqual(len(document.tables), 8)
            document_text = _document_table_text(document)
            self.assertNotIn("\ufffd", document_text)
            self.assertIn("车载摄像头模组", document_text)

    def test_fastapi_generate_and_download_artifacts(self) -> None:
        try:
            from fastapi.testclient import TestClient
        except ImportError as exc:  # pragma: no cover - depends on local test extras.
            self.skipTest(f"FastAPI test client unavailable: {exc}")

        with tempfile.TemporaryDirectory() as directory:
            app = create_sop_fastapi_app(default_out_dir=Path(directory))
            client = TestClient(app)

            health = client.get("/api/sop/health")
            self.assertEqual(health.status_code, 200)
            self.assertEqual(health.json()["draft_status"], "demo_not_for_release")

            payload = _sample_request(Path(directory)).model_dump(mode="json")
            response = client.post("/api/sop/generate", json=payload)
            self.assertEqual(response.status_code, 200, response.text)
            body = response.json()
            self.assertEqual(body["run_id"], "sop_api_test")
            self.assertEqual(body["validation"]["top_level_tables"], 8)
            self.assertFalse(body["validation"]["contains_replacement_char"])

            manifest = client.get("/api/sop/runs/sop_api_test")
            self.assertEqual(manifest.status_code, 200)
            self.assertEqual(manifest.json()["manifest"]["run_id"], "sop_api_test")

            docx = client.get("/api/sop/runs/sop_api_test/artifacts/document_docx")
            self.assertEqual(docx.status_code, 200)
            self.assertGreater(len(docx.content), 1000)
            self.assertIn("application/vnd.openxmlformats-officedocument", docx.headers["content-type"])

            png = client.get("/api/sop/runs/sop_api_test/artifacts/center_flowchart_png")
            self.assertEqual(png.status_code, 200)
            self.assertEqual(png.headers["content-type"], "image/png")
            self.assertGreater(len(png.content), 1000)


def _sample_request(out_dir: Path) -> SopGenerateRequest:
    return SopGenerateRequest(
        product_name="车载摄像头模组",
        part_no="CAM-API-DEMO-001",
        document_no="SOP-CAM-API-001",
        drawing_no="DWG-CAM-API-001",
        station="无尘装配及EOL测试",
        requirement_text="来料核对 -> PCBA装入前壳 -> 镜头组件装配 -> 点胶固化 -> 气密检查 -> EOL测试 -> 合格判定 -> 外观清洁贴标包装",
        routing_steps=[
            SopRoutingStep(name="来料核对", type="inspection", visual_type="inspection", machine_model="IQC-BENCH-DEMO", average_observed_time_s=8),
            SopRoutingStep(name="PCBA装入前壳", type="process", visual_type="assembly", machine_model="ESD-BENCH-DEMO", average_observed_time_s=12),
            SopRoutingStep(name="镜头组件装配", type="process", visual_type="assembly", machine_model="ASSY-JIG-DEMO", average_observed_time_s=14),
            SopRoutingStep(name="点胶固化", type="process", visual_type="process", machine_model="DISPENSE-UV-DEMO", average_observed_time_s=20),
            SopRoutingStep(name="EOL测试", type="test", visual_type="test", machine_model="EOL-TEST-DEMO", average_observed_time_s=18),
            SopRoutingStep(name="外观清洁贴标包装", type="process", visual_type="pack", machine_model="PACK-BENCH-DEMO", average_observed_time_s=10),
        ],
        machine_hints=["IQC-BENCH-DEMO", "ESD-BENCH-DEMO", "ASSY-JIG-DEMO", "DISPENSE-UV-DEMO", "EOL-TEST-DEMO", "PACK-BENCH-DEMO"],
        run_id="sop_api_test",
        out_dir=out_dir,
    )


def _document_table_text(document: Document) -> str:
    values: list[str] = []
    seen: set[object] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if cell.text.strip():
                    values.append(cell.text.strip())
    return "\n".join(values)


if __name__ == "__main__":
    unittest.main()
