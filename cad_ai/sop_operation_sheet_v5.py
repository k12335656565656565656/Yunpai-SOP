from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .sop_visual_template import (
    _fill_word_image_cell,
    _normalize_font_profile,
    _normalize_work_image_slots,
    _set_row_height,
    _set_table_borders,
    _set_word_cell,
    _set_word_cell_margins,
    _set_word_cell_shading,
    _set_word_cell_width,
)


OPERATION_SHEET_LAYOUT_MODE = (
    "portrait_flow_then_excel_reference_operation_sheet_v5"
)

_COLUMN_WIDTHS_CM = [1.1] + [1.4] * 13 + [1.74] * 5
_LEFT_LAST_COLUMN = 13
_RIGHT_FIRST_COLUMN = 14
_RIGHT_LAST_COLUMN = 18


def render_operation_sheet_v5(document: Any, page: dict[str, Any]) -> None:
    """Render one route step using the supplied factory Excel SOP structure."""
    _render_header(document, page)
    _render_operation_area(document, page)
    _render_quality_area(document, page)
    _render_signoff(document)


def _render_header(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=3, cols=19)
    _prepare_table(table)
    for row in table.rows:
        _set_row_height(row, 360)

    _set_word_cell(
        table.cell(0, 0).merge(table.cell(2, 2)),
        "唯格电子",
        bold=True,
        size=15,
    )
    _set_word_cell(
        table.cell(0, 3).merge(table.cell(2, 4)),
        "系列\nSERIES",
        bold=True,
        size=8.5,
        color="0000CC",
    )
    _set_word_cell(
        table.cell(0, 5).merge(table.cell(2, 8)),
        str(page.get("product_name") or ""),
        bold=True,
        size=10.5,
        color="D40000",
    )
    title = table.cell(0, 9).merge(table.cell(2, 13))
    _set_word_cell(
        title,
        "标准作业指导书\nSTANDARD OPERATION PROCEDURE",
        bold=True,
        size=14,
        color="D000D0",
    )
    title_runs = [run for run in title.paragraphs[0].runs if run.text.strip()]
    if len(title_runs) > 1:
        title_runs[-1].font.size = Pt(8)

    values = (
        ("文件编号", page.get("document_no")),
        ("版本", "DRAFT"),
        ("制作日期", page.get("document_date")),
    )
    for row_index, (label, value) in enumerate(values):
        _set_word_cell(
            table.cell(row_index, 14).merge(table.cell(row_index, 16)),
            label,
            size=8.5,
        )
        _set_word_cell(
            table.cell(row_index, 17).merge(table.cell(row_index, 18)),
            str(value or ""),
            size=8.5,
        )
    _compact_table_cells(table)


def _render_operation_area(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=18, cols=19)
    _prepare_table(table)
    for index, row in enumerate(table.rows):
        height = 270
        if index in (6, 7, 8, 9, 10, 12, 13, 14, 15, 16):
            height = 260
        _set_row_height(row, height)

    profile = _normalize_font_profile(page.get("font_profile"))
    body_size = {"standard": 8.5, "clear_large": 9.5, "large": 10.5}[profile]
    compact_size = {"standard": 7.5, "clear_large": 8.1, "large": 8.8}[profile]
    sections = page.get("operation_sections") or {}

    _heading(
        table.cell(0, 0).merge(table.cell(0, _LEFT_LAST_COLUMN)),
        "作业方法（Operating method）",
        size=9.5,
    )
    _heading(
        table.cell(0, _RIGHT_FIRST_COLUMN).merge(table.cell(0, _RIGHT_LAST_COLUMN)),
        "工程序号（Project No）",
        size=8.5,
    )

    method_lines = []
    if str(sections.get("action") or "").strip():
        method_lines.append(f"工序动作：{sections['action']}")
    if str(sections.get("why") or "").strip():
        method_lines.append(f"作业目的：{sections['why']}")
    method_lines.extend(
        f"{index}. {item}"
        for index, item in enumerate(sections.get("methods") or [], start=1)
        if str(item).strip()
    )
    _set_word_cell(
        table.cell(1, 0).merge(table.cell(5, _LEFT_LAST_COLUMN)),
        "\n".join(method_lines) or "待补充作业方法",
        size=body_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    _right_value(
        table,
        1,
        str(sections.get("step_code") or page.get("page_no") or ""),
        compact_size,
    )
    _right_heading(table, 2, "工程名称（Project name）")
    _right_value(table, 3, str(page.get("station") or "待确认"), compact_size)
    _right_heading(table, 4, "使用材料（Use material）")
    _right_value(
        table,
        (5, 7),
        _numbered_lines((sections.get("materials") or []) + (sections.get("inputs") or [])),
        compact_size,
    )
    _right_heading(table, 8, "使用工治具（The use of tooling）")
    _right_value(
        table,
        (9, 11),
        _numbered_lines((sections.get("equipment") or []) + (sections.get("fixtures") or [])),
        compact_size,
    )
    _right_heading(table, 12, "工艺参数（Process parameter）")
    _right_value(
        table,
        13,
        _numbered_lines(sections.get("parameters") or []),
        compact_size,
    )
    _right_heading(table, 14, "IE 工时（人工填写）")
    _right_value(
        table,
        (15, 17),
        "\n".join(str(item) for item in sections.get("ie_lines") or [])
        or "未提供人工 IE 数据",
        compact_size,
    )

    _render_work_images(table, page, body_size=body_size)
    _compact_table_cells(table)


def _render_work_images(table: Any, page: dict[str, Any], *, body_size: float) -> None:
    slot_count = _normalize_work_image_slots(
        page.get("work_image_slots") or len(page.get("step_slots") or []) or 3
    )
    slots = {
        int(item.get("slot_no") or 0): item for item in page.get("step_slots") or []
    }
    for slot_no, image_rows, caption_rows, columns in _image_layout(slot_count):
        image_cell = table.cell(image_rows[0], columns[0]).merge(
            table.cell(image_rows[1], columns[1])
        )
        caption_cell = table.cell(caption_rows[0], columns[0]).merge(
            table.cell(caption_rows[1], columns[1])
        )
        slot = slots.get(slot_no) or {"slot_no": slot_no, "text_placeholder": ""}
        image_path_text = str(slot.get("image_path") or "").strip()
        image_path = Path(image_path_text) if image_path_text else None
        width_cm = sum(_COLUMN_WIDTHS_CM[columns[0] : columns[1] + 1]) - 0.25
        if image_path is not None and image_path.is_file():
            _fill_word_image_cell(
                image_cell,
                image_path,
                slot_no,
                max_width_cm=width_cm,
                max_height_cm=5.1 if slot_count <= 3 else 2.2,
            )
        else:
            _set_word_cell(
                image_cell,
                f"图 {slot_no}\n未配图",
                bold=True,
                size=max(9, body_size),
                color="666666",
            )
            _set_word_cell_shading(image_cell, "F4F6F7")
        caption = _display_image_caption(str(slot.get("image_caption") or ""))
        _set_word_cell(
            caption_cell,
            f"图 {slot_no}：{caption}" if caption else f"图 {slot_no}",
            size=max(7.5, body_size - 0.8),
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )


def _render_quality_area(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=5, cols=19)
    _prepare_table(table)
    for row, height in zip(table.rows, (260, 260, 350, 350, 350)):
        _set_row_height(row, height)

    profile = _normalize_font_profile(page.get("font_profile"))
    size = {"standard": 7.5, "clear_large": 8.1, "large": 8.8}[profile]
    sections = page.get("operation_sections") or {}
    _heading(
        table.cell(0, 0).merge(table.cell(0, _LEFT_LAST_COLUMN)),
        "品质管制点（Quality control point）",
        size=9,
    )
    _heading(
        table.cell(0, _RIGHT_FIRST_COLUMN).merge(table.cell(0, _RIGHT_LAST_COLUMN)),
        "注意事项",
        size=9,
    )
    _heading(table.cell(1, 0), "NO.", size=8)
    _heading(table.cell(1, 1).merge(table.cell(1, 2)), "管制点", size=8)
    _heading(table.cell(1, 3).merge(table.cell(1, 10)), "判定基准", size=8)
    _heading(table.cell(1, 11).merge(table.cell(1, 13)), "使用工具", size=8)

    checks = [str(item) for item in sections.get("quality_checks") or []]
    acceptance = [str(item) for item in sections.get("acceptance") or []]
    tools = [str(item) for item in (sections.get("equipment") or []) + (sections.get("fixtures") or [])]
    for row_offset in range(3):
        row_index = row_offset + 2
        _set_word_cell(table.cell(row_index, 0), str(row_offset + 1), size=size)
        _set_word_cell(
            table.cell(row_index, 1).merge(table.cell(row_index, 2)),
            checks[row_offset] if row_offset < len(checks) else "",
            size=size,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
        _set_word_cell(
            table.cell(row_index, 3).merge(table.cell(row_index, 10)),
            acceptance[row_offset] if row_offset < len(acceptance) else "",
            size=size,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
        _set_word_cell(
            table.cell(row_index, 11).merge(table.cell(row_index, 13)),
            tools[row_offset] if row_offset < len(tools) else "",
            size=size,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )

    caution = (
        [str(item) for item in sections.get("safety") or []]
        + [str(item) for item in sections.get("exceptions") or []]
        + [str(item) for item in sections.get("records") or []]
    )
    _set_word_cell(
        table.cell(1, _RIGHT_FIRST_COLUMN).merge(table.cell(4, _RIGHT_LAST_COLUMN)),
        _numbered_lines(caution),
        size=size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _compact_table_cells(table)


def _render_signoff(document: Any) -> None:
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _set_row_height(table.rows[0], 300)
    for column, label in enumerate(("核准：", "审核：", "制表：")):
        _set_word_cell(
            table.cell(0, column),
            label,
            size=9,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
    _compact_table_cells(table)


def _prepare_table(table: Any) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for column, width in enumerate(_COLUMN_WIDTHS_CM):
        table.columns[column].width = Cm(width)
        for row in table.rows:
            _set_word_cell_width(row.cells[column], width)


def _right_heading(table: Any, row: int, text: str) -> None:
    _heading(
        table.cell(row, _RIGHT_FIRST_COLUMN).merge(table.cell(row, _RIGHT_LAST_COLUMN)),
        text,
        size=8,
    )


def _right_value(table: Any, rows: int | tuple[int, int], text: str, size: float) -> None:
    start, end = (rows, rows) if isinstance(rows, int) else rows
    cell = table.cell(start, _RIGHT_FIRST_COLUMN).merge(table.cell(end, _RIGHT_LAST_COLUMN))
    _set_word_cell(cell, text or "待确认", size=size, align=WD_ALIGN_PARAGRAPH.LEFT)


def _heading(cell: Any, text: str, *, size: float) -> None:
    _set_word_cell(cell, text, bold=True, size=size, shaded=True)


def _numbered_lines(values: list[Any]) -> str:
    lines = [str(item).strip() for item in values if str(item).strip()]
    return "\n".join(f"{index}. {item}" for index, item in enumerate(lines, start=1)) or "待确认"


def _image_layout(
    slot_count: int,
) -> list[tuple[int, tuple[int, int], tuple[int, int], tuple[int, int]]]:
    count = _normalize_work_image_slots(slot_count)
    if count <= 3:
        return [
            (index + 1, (6, 15), (16, 17), columns)
            for index, columns in enumerate(_column_ranges(count))
        ]
    top_count = 2 if count == 4 else 3
    bottom_count = count - top_count
    result = [
        (index + 1, (6, 10), (11, 11), columns)
        for index, columns in enumerate(_column_ranges(top_count))
    ]
    result.extend(
        (top_count + index + 1, (12, 16), (17, 17), columns)
        for index, columns in enumerate(_column_ranges(bottom_count))
    )
    return result


def _column_ranges(count: int) -> list[tuple[int, int]]:
    return [
        (
            round(index * 14 / count),
            round((index + 1) * 14 / count) - 1,
        )
        for index in range(count)
    ]


def _display_image_caption(value: str) -> str:
    caption = value.strip()
    stem = Path(caption).stem
    if re.fullmatch(r"[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}", stem):
        return ""
    return caption


def _compact_table_cells(table: Any) -> None:
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            _remove_merge_artifact_paragraphs(cell)
            _set_word_cell_margins(cell, top=20, start=45, bottom=20, end=45)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)


def _remove_merge_artifact_paragraphs(cell: Any) -> None:
    """Drop empty paragraphs retained when python-docx vertically merges cells."""
    paragraphs = list(cell.paragraphs)
    if len(paragraphs) <= 1:
        return
    for paragraph in paragraphs[1:]:
        element = paragraph._element
        if paragraph.text or element.xpath(".//w:drawing") or element.xpath(".//w:pict"):
            continue
        element.getparent().remove(element)
