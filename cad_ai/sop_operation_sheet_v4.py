from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .sop_visual_template import (
    _clear_word_cell,
    _fill_word_image_cell,
    _normalize_font_profile,
    _normalize_work_image_slots,
    _set_row_height,
    _set_table_borders,
    _set_word_cell,
    _set_word_cell_margins,
    _set_word_cell_shading,
    _set_word_cell_width,
)


OPERATION_SHEET_LAYOUT_MODE = (
    "portrait_flow_then_repeated_landscape_operation_sheet_v4"
)


def render_operation_sheet_v4(document: Any, page: dict[str, Any]) -> None:
    """Render one route step as a readable, editable landscape operation sheet."""
    _render_header(document, page)
    _render_main_content(document, page)
    _render_quality_content(document, page)
    _render_footer(document, page)


def _render_header(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=3, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for row, height in zip(table.rows, (500, 420, 440)):
        _set_row_height(row, height)

    _set_word_cell(table.cell(0, 0), "LOGO", bold=True, size=13, color="F28C00")
    _set_word_cell(
        table.cell(0, 1).merge(table.cell(0, 6)),
        "标准作业指导书",
        bold=True,
        size=18,
    )
    _set_word_cell(
        table.cell(0, 7),
        f"页码\n{page.get('page_no')} OF {page.get('page_total')}",
        size=8,
    )
    values = (
        ("产品品名", page.get("product_name")),
        ("本厂料号", page.get("part_no")),
        ("文件编号", page.get("document_no")),
        ("制作日期", page.get("document_date")),
    )
    for index, (label, value) in enumerate(values):
        _set_word_cell(table.cell(1, index * 2), label, shaded=True, bold=True, size=8)
        _set_word_cell(table.cell(1, index * 2 + 1), str(value or ""), size=8)
    _set_word_cell(table.cell(2, 0), "工站", shaded=True, bold=True, size=8)
    _set_word_cell(table.cell(2, 1), str(page.get("station") or ""), size=9)
    _set_word_cell(table.cell(2, 2), "版本", shaded=True, bold=True, size=8)
    _set_word_cell(table.cell(2, 3), "DRAFT", bold=True, size=9)
    _set_word_cell(table.cell(2, 4), "作业顺序", shaded=True, bold=True, size=8)
    _set_word_cell(
        table.cell(2, 5).merge(table.cell(2, 7)),
        str(page.get("operation_order") or ""),
        size=8.5,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _compact_table_cells(table)


def _render_main_content(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=10, cols=10)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = [2.85] * 6 + [2.55] * 4
    for column, width in enumerate(widths):
        table.columns[column].width = Cm(width)
        for row in table.rows:
            _set_word_cell_width(row.cells[column], width)
    for row, height in zip(
        table.rows,
        (300, 520, 520, 300, 620, 620, 380, 620, 620, 380),
    ):
        _set_row_height(row, height)

    profile = _normalize_font_profile(page.get("font_profile"))
    body_size = {"standard": 8.5, "clear_large": 9.5, "large": 10.5}[profile]
    compact_size = {"standard": 7.5, "clear_large": 8.25, "large": 9.0}[profile]
    sections = page.get("operation_sections") or {}

    _section_title(table.cell(0, 0).merge(table.cell(0, 5)), "作业方法")
    _section_title(table.cell(0, 6).merge(table.cell(0, 9)), "工序信息")
    method_cell = table.cell(1, 0).merge(table.cell(2, 5))
    method_lines = [
        f"工序动作：{sections.get('action') or '待确认'}",
        f"作业目的：{sections.get('why') or '待确认'}",
    ]
    method_lines.extend(
        f"{index}. {item}"
        for index, item in enumerate(sections.get("methods") or [], start=1)
        if str(item).strip()
    )
    _set_word_cell(
        method_cell,
        "\n".join(method_lines),
        size=body_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    _set_word_cell(
        table.cell(1, 6).merge(table.cell(1, 9)),
        "\n".join(
            [
                f"工序编号：{sections.get('step_code') or '待确认'}",
                f"工序名称：{page.get('station') or '待确认'}",
            ]
        ),
        size=compact_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_word_cell(
        table.cell(2, 6).merge(table.cell(2, 9)),
        _labelled_lines(
            ("材料与输入", sections.get("materials")),
            ("输入资料", sections.get("inputs")),
        ),
        size=compact_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _section_title(table.cell(3, 0).merge(table.cell(3, 5)), "作业图片")
    _set_word_cell(
        table.cell(3, 6).merge(table.cell(3, 9)),
        _labelled_lines(
            ("设备/工具", sections.get("equipment")),
            ("治具", sections.get("fixtures")),
        ),
        size=compact_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_word_cell(
        table.cell(4, 6).merge(table.cell(4, 9)),
        _labelled_lines(("工艺参数", sections.get("parameters"))),
        size=compact_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _section_title(table.cell(5, 6).merge(table.cell(5, 9)), "IE 工时（仅显示人工填写值）")
    _set_word_cell(
        table.cell(6, 6).merge(table.cell(9, 9)),
        "\n".join(sections.get("ie_lines") or ["未提供人工 IE 数据"]),
        size=compact_size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _render_work_images(table, page, body_size=body_size)
    _compact_table_cells(table)


def _render_work_images(table: Any, page: dict[str, Any], *, body_size: float) -> None:
    slot_count = _normalize_work_image_slots(
        page.get("work_image_slots") or len(page.get("step_slots") or []) or 3
    )
    slots = {
        int(item.get("slot_no") or 0): item for item in page.get("step_slots") or []
    }
    for slot_no, image_rows, caption_rows, columns in _image_layout(slot_count):
        image_cell = table.cell(image_rows[0], columns[0]).merge(
            table.cell(image_rows[1], columns[1])
        )
        caption_cell = table.cell(caption_rows[0], columns[0]).merge(
            table.cell(caption_rows[1], columns[1])
        )
        slot = slots.get(slot_no) or {"slot_no": slot_no, "text_placeholder": ""}
        image_path_text = str(slot.get("image_path") or "").strip()
        image_path = Path(image_path_text) if image_path_text else None
        if image_path is not None and image_path.is_file():
            _fill_word_image_cell(
                image_cell,
                image_path,
                slot_no,
                max_width_cm=(columns[1] - columns[0] + 1) * 2.85 - 0.35,
                max_height_cm=3.45 if slot_count <= 3 else 1.55,
            )
        else:
            _set_word_cell(
                image_cell,
                f"图 {slot_no}\n未配图",
                bold=True,
                size=max(9, body_size),
                color="666666",
            )
            _set_word_cell_shading(image_cell, "F2F4F5")
        caption = _clean_caption(str(slot.get("text_placeholder") or ""), slot_no)
        _set_word_cell(
            caption_cell,
            f"{slot_no}. {caption}" if caption else f"{slot_no}. 待补充作业说明",
            size=max(7.5, body_size - 0.5),
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )


def _render_quality_content(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=3, cols=10)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for row, height in zip(table.rows, (280, 300, 860)):
        _set_row_height(row, height)
    profile = _normalize_font_profile(page.get("font_profile"))
    size = {"standard": 7.5, "clear_large": 8.25, "large": 9.0}[profile]
    sections = page.get("operation_sections") or {}
    _section_title(table.cell(0, 0).merge(table.cell(0, 9)), "品质、注意与记录")
    items = (
        ("检查方法", sections.get("quality_checks")),
        ("合格判据", sections.get("acceptance")),
        ("检查工具/治具", (sections.get("equipment") or []) + (sections.get("fixtures") or [])),
        ("安全要求", sections.get("safety")),
        (
            "异常处理/记录",
            (sections.get("exceptions") or []) + (sections.get("records") or []),
        ),
    )
    for index, (label, values) in enumerate(items):
        start = index * 2
        _section_title(table.cell(1, start).merge(table.cell(1, start + 1)), label, size=8)
        _set_word_cell(
            table.cell(2, start).merge(table.cell(2, start + 1)),
            "\n".join(str(item) for item in values or []) or "待确认",
            size=size,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
    _compact_table_cells(table)


def _render_footer(document: Any, page: dict[str, Any]) -> None:
    table = document.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for row, height in zip(table.rows, (220, 260)):
        _set_row_height(row, height)
    labels = ["批准", "审核", "制作", "材料环保要求", "管制文件（印章处）", "图号"]
    for column, label in enumerate(labels):
        _set_word_cell(table.cell(0, column), label, shaded=True, bold=True, size=7.5)
    for column in range(3):
        _set_word_cell(table.cell(1, column), "")
    _set_word_cell(
        table.cell(1, 3),
        "材料符合 RoHS/REACH；发布前确认。",
        size=7,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_word_cell(table.cell(1, 4), "")
    _set_word_cell(
        table.cell(1, 5),
        str(page.get("drawing_no") or ""),
        size=7,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _compact_table_cells(table)


def _image_layout(slot_count: int) -> list[tuple[int, tuple[int, int], tuple[int, int], tuple[int, int]]]:
    count = _normalize_work_image_slots(slot_count)
    column_sets = {
        1: [(0, 5)],
        2: [(0, 2), (3, 5)],
        3: [(0, 1), (2, 3), (4, 5)],
    }
    if count <= 3:
        return [
            (index + 1, (4, 7), (8, 9), columns)
            for index, columns in enumerate(column_sets[count])
        ]
    top_count = 2 if count == 4 else 3
    bottom_count = count - top_count
    result = [
        (index + 1, (4, 5), (6, 6), columns)
        for index, columns in enumerate(column_sets[top_count])
    ]
    result.extend(
        (top_count + index + 1, (7, 8), (9, 9), columns)
        for index, columns in enumerate(column_sets[bottom_count])
    )
    return result


def _section_title(cell: Any, text: str, *, size: float = 8.5) -> None:
    _set_word_cell(cell, text, bold=True, size=size, shaded=True)


def _labelled_lines(*items: tuple[str, Any]) -> str:
    lines: list[str] = []
    for label, values in items:
        normalized = [str(item).strip() for item in values or [] if str(item).strip()]
        lines.append(f"{label}：{'；'.join(normalized) if normalized else '待确认'}")
    return "\n".join(lines)


def _clean_caption(value: str, slot_no: int) -> str:
    text = value.strip()
    prefix = f"{slot_no}."
    return text[len(prefix):].strip() if text.startswith(prefix) else text


def _compact_table_cells(table: Any) -> None:
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            _set_word_cell_margins(cell, top=25, start=55, bottom=25, end=55)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
