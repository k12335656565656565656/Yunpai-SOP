from __future__ import annotations

import tempfile
import unittest
import json
import os
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import patch

from cad_ai.sop_knowledge.conversation import SopConversationService
from cad_ai.sop_knowledge.documents import (
    CURRENT_PREVIEW_DIR_NAME,
    MULTI_PAGE_TEMPLATE_ID,
    VERSIONED_PREVIEW_DIR_PREFIX,
    SopDocumentService,
)
from cad_ai.sop_knowledge.models import RouteSectionDraft
from cad_ai.sop_knowledge.nl_assistant import NaturalLanguageSopAssistant
from cad_ai.sop_knowledge.store import SopKnowledgeStore
from tests.test_sop_knowledge_workflow import make_identity, make_route


class FakeAssistant:
    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        step = route["steps"][1]
        return ({
            "assistant_message": "我定位到第二道工序和质量控制章节，已经按描述修改草稿。",
            "judgement": ["“第二步”对应当前路线排序中的第二道工序。"],
            "summary": "修改 1 个工序字段和 1 个章节。",
            "changes": [{
                "step_id": step["id"], "step_code": step["step_code"], "step_title": step["title"],
                "field_name": "method", "value": ["先核对方向", "插入到位", "轻拉确认"],
                "reason": "用户要求拆解动作",
            }],
            "new_steps": [],
            "section_changes": [{
                "section_type": "quality_control", "patch": {"inspection_note": "逐件轻拉确认"},
                "reason": "用户补充检查要求",
            }],
            "image_refs": [], "warnings": [], "requires_human_confirmation": True,
        }, "llm")


class UnsafeFallbackNewStepAssistant:
    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        return ({
            "assistant_message": "已识别到两个待新增工序。",
            "judgement": [],
            "summary": "新增两个工序。",
            "changes": [],
            "new_steps": [
                {"title": "3我觉得不够齐全", "method": []},
                {"title": "需要补充一下", "method": []},
            ],
            "section_changes": [],
            "image_refs": [],
            "warnings": ["AI 服务暂时不可用，已使用离线规则解析。"],
            "requires_human_confirmation": True,
        }, "deterministic_fallback")


class LockedTargetAssistant:
    def __init__(self) -> None:
        self.calls = 0

    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        self.calls += 1
        step_id = int(route["_locked_target_step_id"])
        step = next(item for item in route["steps"] if int(item["id"]) == step_id)
        return ({
            "assistant_message": f"已定位到{step['title']}，修改内容已写入草稿。",
            "judgement": [f"本次只修改{step['title']}。"],
            "summary": "修改 1 项安全要求。",
            "changes": [{
                "step_id": step["id"], "step_code": step["step_code"], "step_title": step["title"],
                "field_name": "safety", "value": ["操作前确认设备状态"],
                "reason": "用户明确补充安全要求",
            }],
            "new_steps": [], "section_changes": [], "image_refs": [], "warnings": [],
            "requires_human_confirmation": True,
        }, "llm")


class WrongTargetAssistant:
    def __init__(self) -> None:
        self.calls = 0

    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        self.calls += 1
        step = route["steps"][0]
        return ({
            "assistant_message": "已完成修改。",
            "judgement": "已定位工序。",
            "summary": "修改 1 项安全要求。",
            "changes": [{
                "step_id": step["id"], "step_code": step["step_code"], "step_title": step["title"],
                "field_name": "safety", "value": ["操作前确认设备状态"],
                "reason": "错误沿用了旧目标",
            }],
            "new_steps": [], "section_changes": [], "image_refs": [], "warnings": [],
            "requires_human_confirmation": True,
        }, "llm")


class NeverCalledAssistant:
    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        raise AssertionError("ambiguous target must be confirmed before calling the assistant")


class ImageLayoutOperationAssistant:
    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        step_id = int(route.get("_locked_target_step_id") or route["steps"][1]["id"])
        step = next(item for item in route["steps"] if int(item["id"]) == step_id)
        return ({
            "assistant_message": "已识别图片格数调整。",
            "judgement": ["图片版式调整需要人工确认。"],
            "summary": "将工图改为 2 格。",
            "changes": [], "new_steps": [], "section_changes": [], "image_refs": [],
            "operations": [{
                "kind": "set_image_slots", "step_ref": str(step["id"]), "step_id": step["id"],
                "step_code": step["step_code"], "step_title": step["title"], "slots": 2,
                "reason": "用户要求图片更清晰。",
            }],
            "warnings": [], "requires_human_confirmation": True,
        }, "llm")


class PreviewNavigationAssistant:
    def preview(self, instruction: str, route: dict[str, Any], *, history=None):
        return ({
            "assistant_message": "已定位到第 3 页。",
            "judgement": [], "summary": "跳转预览。",
            "changes": [], "new_steps": [], "section_changes": [], "image_refs": [],
            "operations": [{"kind": "navigate_preview", "page": 3, "reason": "用户要求查看指定页。"}],
            "warnings": [], "requires_human_confirmation": False,
        }, "llm")


class FakeDocuments:
    def __init__(self) -> None:
        self.generated: list[int] = []

    def generate(self, route_id: int) -> dict[str, Any]:
        self.generated.append(route_id)
        return {
            "route_id": route_id, "route_version": 1, "product_code": "CHAT-TEST",
            "generated_at": "2026-08-12T00:00:00+00:00", "version_token": "test",
            "page_count": 3, "media_count": 0, "status": "draft_document_generated",
            "preview_source": "generated_docx", "docx_url": "/latest.docx",
            "preview_url": "/preview.pdf", "page_urls": [],
        }

    def latest(self, route_id: int, *, generate_if_missing: bool = True) -> dict[str, Any]:
        return self.generate(route_id)


class StableDocuments(FakeDocuments):
    def latest(self, route_id: int, *, generate_if_missing: bool = True) -> dict[str, Any]:
        return {
            "route_id": route_id, "route_version": 1, "product_code": "CHAT-TEST",
            "generated_at": "2026-08-12T00:00:00+00:00", "version_token": "existing",
            "page_count": 4, "media_count": 0, "status": "draft_document_generated",
            "preview_source": "generated_docx", "docx_url": "/latest.docx",
            "preview_url": "/preview.pdf", "page_urls": [],
        }


class SopConversationalDocxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.store = SopKnowledgeStore(Path(self.temp.name) / "knowledge.sqlite3")
        self.store.initialize()
        self.store.ensure_process_family("test_family", "测试工艺族")
        identity = make_identity("CHAT-TEST")
        self.store.upsert_product(identity, {"class": "cable"})
        self.route_id = self.store.create_route(make_route(identity, 3))
        for section_type in (
            "product_identity", "bom_material", "equipment_fixture", "process_parameter",
            "quality_control", "packaging_label", "ie_timing", "release_signoff",
        ):
            self.store.create_route_section(
                self.route_id,
                RouteSectionDraft(section_type=section_type, content={"section": section_type}),
            )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_chat_applies_ai_routing_regenerates_docx_and_records_work(self) -> None:
        documents = FakeDocuments()
        service = SopConversationService(
            self.store, documents, assistant=FakeAssistant()  # type: ignore[arg-type]
        )
        result = service.chat(self.route_id, "把第二步拆详细，并补上质检说明", worker="worker-01")
        self.assertTrue(result["docx_regenerated"])
        self.assertEqual(documents.generated, [self.route_id])
        self.assertEqual(len(result["changes"]), 2)
        route = self.store.get_route(self.route_id)
        self.assertEqual(route["steps"][1]["method_json"], ["先核对方向", "插入到位", "轻拉确认"])
        quality = next(item for item in route["sections"] if item["section_type"] == "quality_control")
        self.assertEqual(quality["version"], 2)
        self.assertEqual(quality["content_json"]["inspection_note"], "逐件轻拉确认")
        history = self.store.list_chat_messages(self.route_id)
        self.assertEqual([item["role"] for item in history], ["user", "assistant"])
        self.assertTrue(history[-1]["metadata_json"]["docx_regenerated"])
        self.assertIn("judgement", history[-1]["metadata_json"])
        self.assertEqual(result["changes"][0]["field"], "作业步骤")
        self.assertEqual(result["changes"][0]["location"], "DOCX 第 3 页 > 作业步骤")
        self.assertEqual(result["changes"][0]["page_number"], 3)
        self.assertEqual(result["changes"][0]["field_key"], "method")

    def test_explicit_read_only_question_never_applies_model_changes(self) -> None:
        documents = FakeDocuments()
        service = SopConversationService(
            self.store, documents, assistant=FakeAssistant()  # type: ignore[arg-type]
        )

        result = service.chat(
            self.route_id,
            "请说明第 2 道工序当前内容，不要修改 SOP。",
            worker="worker-01",
        )

        self.assertFalse(result["docx_regenerated"])
        self.assertEqual(result["changes"], [])
        self.assertEqual(documents.generated, [self.route_id])
        self.assertIn("未写入 SOP 草稿", result["message"])
        self.assertEqual(self.store.get_route(self.route_id)["steps"][1]["review_state"], "unreviewed")

    def test_conversational_layout_waits_for_text_confirmation_then_regenerates(self) -> None:
        documents = StableDocuments()
        service = SopConversationService(
            self.store, documents, assistant=ImageLayoutOperationAssistant()  # type: ignore[arg-type]
        )

        pending = service.chat(self.route_id, "把第 2 道工序调整成 2 格工图", worker="worker-01")

        self.assertEqual(pending["parser_kind"], "operation_confirmation")
        self.assertFalse(pending["docx_regenerated"])
        self.assertIn("确认执行", pending["message"])
        self.assertEqual(self.store.get_route(self.route_id)["steps"][1]["work_image_slots"], 3)

        still_pending = service.chat(self.route_id, "先说明影响", worker="worker-01")
        self.assertEqual(still_pending["parser_kind"], "operation_confirmation")
        self.assertFalse(still_pending["docx_regenerated"])
        self.assertEqual(self.store.get_route(self.route_id)["steps"][1]["work_image_slots"], 3)

        applied = service.chat(self.route_id, "确认执行", worker="worker-01")

        self.assertEqual(applied["parser_kind"], "operation_applied")
        self.assertTrue(applied["docx_regenerated"])
        self.assertEqual(documents.generated, [self.route_id])
        self.assertEqual(self.store.get_route(self.route_id)["steps"][1]["work_image_slots"], 2)
        self.assertEqual(applied["changes"][-1]["field"], "路线操作")

    def test_conversational_preview_navigation_never_mutates_document(self) -> None:
        documents = StableDocuments()
        service = SopConversationService(
            self.store, documents, assistant=PreviewNavigationAssistant()  # type: ignore[arg-type]
        )

        result = service.chat(self.route_id, "跳到第 3 页", worker="worker-01")

        self.assertEqual(result["parser_kind"], "preview_navigation")
        self.assertFalse(result["docx_regenerated"])
        self.assertEqual(result["preview_navigation"], {"page": 3})
        self.assertEqual(documents.generated, [])

    def test_offline_parser_recognizes_clear_route_actions_without_guessing(self) -> None:
        assistant = NaturalLanguageSopAssistant(use_llm=False)
        route = self.store.get_route(self.route_id)

        layout, parser_kind = assistant.preview("把第 2 道工序的工图改成 2 格", route)
        self.assertEqual(parser_kind, "deterministic")
        self.assertEqual(layout["operations"], [{
            "kind": "set_image_slots", "step_ref": str(route["steps"][1]["id"]),
            "step_id": route["steps"][1]["id"], "step_code": route["steps"][1]["step_code"],
            "step_title": route["steps"][1]["title"], "slots": 2, "reason": "用户要求调整指导书图片格数",
        }])

        deletion, _ = assistant.preview("删除第 3 道工序", route)
        self.assertEqual(deletion["operations"][0]["kind"], "delete_step")
        self.assertEqual(deletion["operations"][0]["step_id"], route["steps"][2]["id"])

        split, _ = assistant.preview("把第 2 道工序拆成：校对、加工", route)
        self.assertEqual(split["operations"][0]["kind"], "split_actions")
        self.assertEqual(split["operations"][0]["titles"], ["校对", "加工"])

        merge, _ = assistant.preview("把第 1 道工序和第 2 道工序合并为前段作业", route)
        self.assertEqual(merge["operations"][0]["kind"], "merge_steps")
        self.assertEqual(merge["operations"][0]["step_id"], route["steps"][0]["id"])
        self.assertEqual(merge["operations"][0]["source_steps"][0]["step_id"], route["steps"][1]["id"])

        navigation, _ = assistant.preview("跳到第 3 页", route)
        self.assertEqual(navigation["operations"], [{
            "kind": "navigate_preview", "page": 3, "reason": "用户要求跳转预览页",
        }])

        unclear, _ = assistant.preview("把这个工序拆开", route)
        self.assertEqual(unclear["operations"], [])

    def test_font_profile_choice_waits_for_confirmation_then_regenerates_once(self) -> None:
        documents = StableDocuments()
        service = SopConversationService(
            self.store, documents, assistant=NeverCalledAssistant()  # type: ignore[arg-type]
        )

        prompt = service.chat(self.route_id, "字体调大一点", worker="worker-01")

        self.assertEqual(prompt["parser_kind"], "font_profile_selection")
        self.assertFalse(prompt["docx_regenerated"])
        self.assertEqual(documents.generated, [])
        self.assertIn("1. 标准阅读", prompt["message"])
        self.assertIn("2. 清晰大字", prompt["message"])
        self.assertIn("3. 大字版", prompt["message"])
        self.assertEqual(self.store.get_route(self.route_id)["route"]["font_profile"], "standard")

        result = service.chat(self.route_id, "2", worker="worker-01")

        self.assertEqual(result["parser_kind"], "font_profile")
        self.assertTrue(result["docx_regenerated"])
        self.assertEqual(documents.generated, [self.route_id])
        self.assertEqual(self.store.get_route(self.route_id)["route"]["font_profile"], "clear_large")
        with self.store.connect() as connection:
            decision = connection.execute(
                "SELECT entity_type,field_name,decision FROM review_decision WHERE field_name='font_profile'"
            ).fetchone()
        self.assertEqual(dict(decision), {
            "entity_type": "route", "field_name": "font_profile", "decision": "needs_revision",
        })

        repeat_prompt = service.chat(self.route_id, "字体调大一点", worker="worker-01")
        repeated = service.chat(self.route_id, "2", worker="worker-01")
        self.assertFalse(repeat_prompt["docx_regenerated"])
        self.assertFalse(repeated["docx_regenerated"])
        self.assertEqual(documents.generated, [self.route_id])

    def test_font_profile_on_approved_route_creates_revision_only_after_choice(self) -> None:
        documents = StableDocuments()
        service = SopConversationService(
            self.store, documents, assistant=NeverCalledAssistant()  # type: ignore[arg-type]
        )
        with self.store.connect() as connection:
            connection.execute("UPDATE product_route SET status='approved' WHERE id=?", (self.route_id,))

        prompt = service.chat(self.route_id, "把字体调大一点", worker="worker-01")
        self.assertEqual(prompt["route_id"], self.route_id)
        self.assertFalse(prompt["docx_regenerated"])
        with self.store.connect() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM product_route").fetchone()[0], 1)

        result = service.chat(self.route_id, "3", worker="worker-01")
        revised_route_id = int(result["route_id"])
        self.assertNotEqual(revised_route_id, self.route_id)
        self.assertEqual(self.store.get_route(self.route_id)["route"]["status"], "approved")
        revised = self.store.get_route(revised_route_id)["route"]
        self.assertEqual(revised["status"], "draft")
        self.assertEqual(revised["font_profile"], "large")
        self.assertTrue(result["docx_regenerated"])
        self.assertEqual(documents.generated, [revised_route_id])

    def test_fallback_never_auto_adds_steps_without_explicit_new_step_intent(self) -> None:
        documents = FakeDocuments()
        service = SopConversationService(
            self.store,
            documents,
            assistant=UnsafeFallbackNewStepAssistant(),  # type: ignore[arg-type]
        )
        before_count = len(self.store.get_route(self.route_id)["steps"])

        result = service.chat(
            self.route_id,
            "工序3我觉得不够齐全，需要补充一下",
            worker="worker-01",
        )

        self.assertEqual(result["parser_kind"], "deterministic_fallback")
        self.assertFalse(result["docx_regenerated"])
        self.assertEqual(len(self.store.get_route(self.route_id)["steps"]), before_count)
        self.assertIn("没有明确要求新增工序", " ".join(result["warnings"]))

    def test_offline_explicit_new_step_intent_still_adds_one_reviewable_step(self) -> None:
        documents = FakeDocuments()
        service = SopConversationService(
            self.store,
            documents,
            assistant=NaturalLanguageSopAssistant(use_llm=False),
        )
        before_count = len(self.store.get_route(self.route_id)["steps"])

        result = service.chat(
            self.route_id,
            "新增工序：激光打标",
            worker="worker-01",
        )

        route = self.store.get_route(self.route_id)
        self.assertTrue(result["docx_regenerated"])
        self.assertEqual(len(route["steps"]), before_count + 1)
        self.assertEqual(route["steps"][-1]["title"], "激光打标")
        self.assertEqual(route["steps"][-1]["review_state"], "needs_revision")

    def test_location_query_returns_last_change_docx_page_without_writing(self) -> None:
        documents = FakeDocuments()
        service = SopConversationService(
            self.store, documents, assistant=FakeAssistant()  # type: ignore[arg-type]
        )
        service.chat(self.route_id, "把第 2 道工序补充完整", worker="worker-01")

        result = service.chat(self.route_id, "刚才具体改在哪里查看？", worker="worker-01")

        self.assertEqual(result["parser_kind"], "reference")
        self.assertFalse(result["docx_regenerated"])
        self.assertIn("第 2 道", result["message"])
        self.assertIn("DOCX 第 3 页 > 作业步骤", result["message"])

    def test_ambiguous_target_returns_candidates_without_writing_or_regenerating(self) -> None:
        steps = self.store.get_route(self.route_id)["steps"]
        self.store.update_step_field(steps[0]["id"], "title", "电气性能检验", reviewer="setup")
        self.store.update_step_field(steps[1]["id"], "title", "成品外观检验", reviewer="setup")
        documents = StableDocuments()
        service = SopConversationService(
            self.store, documents, assistant=NeverCalledAssistant()  # type: ignore[arg-type]
        )

        result = service.chat(
            self.route_id,
            "把检验工序的记录要求补充为登记工单号。",
            worker="worker-01",
        )

        self.assertEqual(result["target_resolution"]["status"], "needs_choice")
        self.assertEqual(len(result["target_resolution"]["candidates"]), 2)
        self.assertFalse(result["docx_regenerated"])
        self.assertEqual(documents.generated, [])
        self.assertIsNone(result["proposal_id"])
        history = self.store.list_chat_messages(self.route_id)
        self.assertEqual(history[-1]["metadata_json"]["pending_instruction"], "把检验工序的记录要求补充为登记工单号。")

    def test_candidate_selection_applies_only_to_selected_step(self) -> None:
        steps = self.store.get_route(self.route_id)["steps"]
        self.store.update_step_field(steps[0]["id"], "title", "电气性能检验", reviewer="setup")
        self.store.update_step_field(steps[1]["id"], "title", "成品外观检验", reviewer="setup")
        documents = StableDocuments()
        assistant = LockedTargetAssistant()
        service = SopConversationService(self.store, documents, assistant=assistant)  # type: ignore[arg-type]
        pending = service.chat(
            self.route_id,
            "把检验工序的安全要求补充为操作前确认设备状态。",
            worker="worker-01",
        )

        result = service.chat(
            self.route_id,
            "选择成品外观检验",
            worker="worker-01",
            selected_step_id=steps[1]["id"],
            pending_message_id=pending["assistant_message_id"],
        )

        route = self.store.get_route(self.route_id)
        self.assertTrue(result["docx_regenerated"])
        self.assertEqual(route["steps"][1]["safety_json"], ["操作前确认设备状态"])
        self.assertNotEqual(route["steps"][0]["safety_json"], ["操作前确认设备状态"])
        self.assertEqual(result["target_resolution"]["selected_step_id"], steps[1]["id"])
        self.assertEqual(assistant.calls, 1)

    def test_model_target_outside_locked_step_is_retried_then_blocked(self) -> None:
        steps = self.store.get_route(self.route_id)["steps"]
        self.store.update_step_field(steps[0]["id"], "title", "裁线与长度补偿", reviewer="setup")
        self.store.update_step_field(steps[1]["id"], "title", "异常隔离与人工放行", reviewer="setup")
        documents = StableDocuments()
        assistant = WrongTargetAssistant()
        service = SopConversationService(self.store, documents, assistant=assistant)  # type: ignore[arg-type]

        result = service.chat(
            self.route_id,
            "把隔离步骤的安全要求补充为操作前确认设备状态。",
            worker="worker-01",
        )

        route = self.store.get_route(self.route_id)
        self.assertFalse(result["docx_regenerated"])
        self.assertEqual(documents.generated, [])
        self.assertEqual(assistant.calls, 2)
        self.assertNotEqual(route["steps"][0]["safety_json"], ["操作前确认设备状态"])
        self.assertIn("目标不一致", " ".join(result["warnings"]))

    def test_document_fingerprint_detects_backend_route_edits(self) -> None:
        documents = SopDocumentService(self.store)
        before = documents._route_fingerprint(self.route_id)
        step = self.store.get_route(self.route_id)["steps"][0]
        self.store.update_step_field(
            step["id"], "method", ["人工通过受控接口修改后的新动作"], reviewer="worker-02",
            decision="needs_revision", comment="测试文档失效检测",
        )
        after = documents._route_fingerprint(self.route_id)
        self.assertNotEqual(before, after)

    def test_document_fingerprint_detects_font_profile_change(self) -> None:
        documents = SopDocumentService(self.store)
        before = documents._route_fingerprint(self.route_id)
        self.store.set_route_font_profile(self.route_id, "clear_large", reviewer="worker-02")
        after = documents._route_fingerprint(self.route_id)
        self.assertNotEqual(before, after)

    def test_linux_preview_conversion_uses_libreoffice_backend(self) -> None:
        documents = SopDocumentService(self.store)
        docx_path = Path(self.temp.name) / "source.docx"
        output_dir = Path(self.temp.name) / "preview"
        expected_pdf = output_dir / "source.pdf"

        with (
            patch("cad_ai.sop_knowledge.documents.sys.platform", "linux"),
            patch.object(
                documents,
                "_convert_docx_with_libreoffice",
                return_value=expected_pdf,
            ) as convert_libreoffice,
            patch.object(documents, "_convert_docx_with_windows") as convert_windows,
        ):
            result = documents._convert_docx_to_pdf(docx_path, output_dir)

        self.assertEqual(result, expected_pdf)
        convert_libreoffice.assert_called_once_with(docx_path, output_dir)
        convert_windows.assert_not_called()

    def test_document_runtime_status_reports_server_converter(self) -> None:
        documents = SopDocumentService(self.store)

        with (
            patch("cad_ai.sop_knowledge.documents.sys.platform", "linux"),
            patch.object(
                documents,
                "_find_libreoffice_executable",
                return_value=Path("/usr/bin/libreoffice"),
            ),
        ):
            status = documents.runtime_status()

        self.assertEqual(status["status"], "ready")
        self.assertEqual(status["platform"], "linux")
        self.assertEqual(status["docx_converter"]["backend"], "LibreOffice")
        self.assertTrue(status["pdf_renderer"]["available"])
        self.assertTrue(status["storage"]["available"])

    def test_real_libreoffice_backend_converts_docx_when_available(self) -> None:
        from docx import Document
        from PIL import Image

        documents = SopDocumentService(self.store)
        executable = documents._find_libreoffice_executable()
        if executable is None:
            self.skipTest("LibreOffice is not installed in this environment")

        fixture_dir = Path(self.temp.name) / "服务器预览"
        output_dir = fixture_dir / "转换结果"
        fixture_dir.mkdir(parents=True)
        docx_path = fixture_dir / "预览检查.docx"
        document = Document()
        document.add_heading("SOP preview readiness", level=1)
        document.add_paragraph("LibreOffice and PyMuPDF must both produce readable artifacts.")
        document.save(docx_path)

        pdf_path = documents._convert_docx_with_libreoffice(docx_path, output_dir)
        page_paths = documents._render_pdf_pages(pdf_path, output_dir)

        self.assertTrue(pdf_path.is_file())
        self.assertGreater(pdf_path.stat().st_size, 0)
        self.assertEqual(len(page_paths), 1)
        self.assertGreater(page_paths[0].stat().st_size, 0)
        with Image.open(page_paths[0]) as preview_page:
            self.assertGreaterEqual(preview_page.width, 1100)

    def test_preview_failure_keeps_existing_pdf_and_pages(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}" / "preview"
        output_dir.mkdir(parents=True)
        old_pdf = output_dir / "existing.pdf"
        old_page = output_dir / "page-001.png"
        old_pdf.write_bytes(b"existing-pdf")
        old_page.write_bytes(b"existing-page")
        docx_path = documents.root / f"route_{self.route_id}" / "source.docx"
        docx_path.write_bytes(b"new-docx")

        failed = SimpleNamespace(returncode=1, stdout="", stderr="Word COM unavailable")
        with (
            patch("cad_ai.sop_knowledge.documents.sys.platform", "win32"),
            patch.object(
                documents,
                "_windows_powershell_executable",
                return_value=Path(r"C:\Windows\powershell.exe"),
            ),
            patch("cad_ai.sop_knowledge.documents.subprocess.run", return_value=failed),
        ):
            with self.assertRaisesRegex(RuntimeError, "Word COM unavailable"):
                documents._render_preview(docx_path, output_dir)

        self.assertEqual(old_pdf.read_bytes(), b"existing-pdf")
        self.assertEqual(old_page.read_bytes(), b"existing-page")

    def test_preview_publish_uses_versioned_directory_when_current_preview_is_locked(self) -> None:
        documents = SopDocumentService(self.store)
        route_dir = documents.root / f"route_{self.route_id}"
        current_dir = route_dir / CURRENT_PREVIEW_DIR_NAME
        candidate_dir = route_dir / "preview-candidate"
        current_dir.mkdir(parents=True)
        candidate_dir.mkdir()
        (current_dir / "old.pdf").write_bytes(b"old-preview")
        (candidate_dir / "new.pdf").write_bytes(b"new-preview")

        real_replace = os.replace

        def replace_with_locked_current(source: Path, destination: Path) -> None:
            if Path(source) == current_dir:
                raise PermissionError("preview directory is in use")
            real_replace(source, destination)

        with patch("cad_ai.sop_knowledge.documents.os.replace", side_effect=replace_with_locked_current):
            published_dir = documents._publish_preview_directory(candidate_dir, current_dir)

        self.assertTrue(published_dir.name.startswith(VERSIONED_PREVIEW_DIR_PREFIX))
        self.assertEqual((published_dir / "new.pdf").read_bytes(), b"new-preview")
        self.assertEqual((current_dir / "old.pdf").read_bytes(), b"old-preview")

    def test_generate_skips_reopening_an_existing_unreadable_preview_directory(self) -> None:
        documents = SopDocumentService(self.store)
        route_dir = documents.root / f"route_{self.route_id}"
        current_dir = route_dir / CURRENT_PREVIEW_DIR_NAME
        template_dir = route_dir / "template_package"
        current_dir.mkdir(parents=True)
        template_dir.mkdir()
        docx_path = template_dir / "current.docx"
        docx_path.write_bytes(b"current-docx")
        page_count = 3
        versioned_dir = route_dir / f"{VERSIONED_PREVIEW_DIR_PREFIX}test"
        rendered = {
            "document_docx": str(docx_path),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "expected_page_count": page_count,
            "validation_json": str(template_dir / "validation.json"),
        }
        preview = {
            "pdf_path": str(versioned_dir / "preview.pdf"),
            "page_paths": [str(versioned_dir / f"page-{index:03d}.png") for index in range(1, page_count + 1)],
            "page_count": page_count,
        }
        real_mkdir = Path.mkdir

        def mkdir_with_unreadable_current(path: Path, *args, **kwargs) -> None:
            if path == current_dir:
                raise PermissionError("preview directory cannot be reopened")
            real_mkdir(path, *args, **kwargs)

        with (
            patch.object(Path, "mkdir", new=mkdir_with_unreadable_current),
            patch.object(documents, "_generate_template_package", return_value=rendered),
            patch.object(documents, "_render_preview", return_value=preview) as render_preview,
        ):
            result = documents.generate(self.route_id)

        render_preview.assert_called_once_with(docx_path, current_dir, expected_page_count=page_count)
        self.assertEqual(result["page_count"], page_count)
        self.assertIn(VERSIONED_PREVIEW_DIR_PREFIX, json.loads((route_dir / "document_manifest.json").read_text(encoding="utf-8"))["pdf_path"])

    def test_latest_returns_persisted_preview_failure_without_retrying(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}"
        output_dir.mkdir(parents=True)
        docx_path = output_dir / "current.docx"
        docx_path.write_bytes(b"current-docx")
        rendered = {
            "document_docx": str(docx_path),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "expected_page_count": 3,
            "validation_json": str(output_dir / "validation.json"),
        }

        with (
            patch.object(documents, "_generate_template_package", return_value=rendered) as generate_template,
            patch.object(
                documents,
                "_render_preview",
                side_effect=RuntimeError("DOCX 已生成，但预览转换失败：Word COM unavailable"),
            ) as render_preview,
        ):
            first = documents.latest(self.route_id)
            second = documents.latest(self.route_id)

        self.assertEqual(generate_template.call_count, 1)
        self.assertEqual(render_preview.call_count, 1)
        self.assertEqual(first["preview_status"], "failed")
        self.assertEqual(second["version_token"], first["version_token"])
        self.assertEqual(first["page_urls"], [])
        self.assertIn("可下载", first["preview_error"])
        resolved, mime_type, _ = documents.resolve_file(self.route_id, "docx")
        self.assertEqual(resolved.resolve(), docx_path.resolve())
        self.assertIn("wordprocessingml", mime_type)

    def test_latest_regenerates_when_manifest_uses_legacy_preview_directory(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}"
        output_dir.mkdir(parents=True)
        docx_path = output_dir / "current.docx"
        pdf_path = output_dir / "preview" / "current.pdf"
        docx_path.write_bytes(b"current-docx")
        pdf_path.parent.mkdir()
        pdf_path.write_bytes(b"current-pdf")
        manifest = {
            "route_id": self.route_id,
            "route_version": 1,
            "product_code": "CHAT-TEST",
            "generated_at": "2026-08-12T00:00:00+00:00",
            "version_token": "unreadable-preview",
            "route_fingerprint": documents._route_fingerprint(self.route_id),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "layout_mode": "portrait_flow_then_repeated_landscape_work_instructions",
            "docx_path": str(docx_path),
            "pdf_path": str(pdf_path),
            "page_paths": [],
            "page_count": 0,
            "expected_page_count": 1,
            "validation_path": "",
            "media_count": 0,
            "status": "draft_document_generated",
            "preview_source": "generated_docx",
        }
        (output_dir / "document_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        regenerated = {"version_token": "rebuilt"}

        with patch.object(documents, "generate", return_value=regenerated) as generate:
            result = documents.latest(self.route_id)

        self.assertEqual(result, regenerated)
        generate.assert_called_once_with(self.route_id)

    def test_latest_regenerates_incomplete_manifest_without_reopening_old_pdf(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}"
        preview_dir = output_dir / CURRENT_PREVIEW_DIR_NAME
        output_dir.mkdir(parents=True)
        preview_dir.mkdir(parents=True)
        docx_path = output_dir / "test.docx"
        pdf_path = preview_dir / "test.pdf"
        docx_path.write_bytes(b"test-docx")
        pdf_path.write_bytes(b"locked-preview-pdf")
        manifest = {
            "route_id": self.route_id,
            "route_version": 1,
            "product_code": "CHAT-TEST",
            "generated_at": "2026-08-12T00:00:00+00:00",
            "version_token": "test",
            "route_fingerprint": documents._route_fingerprint(self.route_id),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "layout_mode": "portrait_flow_then_repeated_landscape_work_instructions",
            "docx_path": str(docx_path),
            "pdf_path": str(pdf_path),
            "page_paths": [],
            "page_count": 1,
            "expected_page_count": 1,
            "validation_path": "",
            "media_count": 0,
            "status": "draft_document_generated",
            "preview_source": "generated_docx",
        }
        (output_dir / "document_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

        regenerated = {"version_token": "fresh-preview"}
        with (
            patch.object(documents, "generate", return_value=regenerated) as generate,
            patch.object(documents, "_render_pdf_pages", side_effect=PermissionError("PDF is locked")) as render_pages,
        ):
            result = documents.latest(self.route_id)

        self.assertEqual(result, regenerated)
        generate.assert_called_once_with(self.route_id)
        render_pages.assert_not_called()

    def test_latest_keeps_readable_page_preview_when_pdf_is_browser_locked(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}"
        preview_dir = output_dir / CURRENT_PREVIEW_DIR_NAME
        output_dir.mkdir(parents=True)
        preview_dir.mkdir(parents=True)
        docx_path = output_dir / "current.docx"
        pdf_path = preview_dir / "current.pdf"
        page_path = preview_dir / "page-001.png"
        docx_path.write_bytes(b"current-docx")
        pdf_path.write_bytes(b"locked-preview-pdf")
        page_path.write_bytes(b"readable-page")
        manifest = {
            "route_id": self.route_id,
            "route_version": 1,
            "product_code": "CHAT-TEST",
            "generated_at": "2026-08-21T00:00:00+00:00",
            "version_token": "locked-pdf",
            "route_fingerprint": documents._route_fingerprint(self.route_id),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "layout_mode": "portrait_flow_then_repeated_landscape_work_instructions",
            "docx_path": str(docx_path),
            "pdf_path": str(pdf_path),
            "page_paths": [str(page_path)],
            "page_count": 1,
            "expected_page_count": 1,
            "validation_path": "",
            "media_count": 0,
            "status": "draft_document_generated",
            "preview_source": "generated_docx",
        }
        (output_dir / "document_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        original_is_file = Path.is_file
        with (
            patch.object(
                Path,
                "is_file",
                autospec=True,
                side_effect=lambda candidate: False if Path(candidate) == pdf_path else original_is_file(candidate),
            ),
            patch.object(documents, "generate") as generate,
            patch.object(documents, "_render_pdf_pages", side_effect=PermissionError("PDF is locked")) as render_pages,
        ):
            result = documents.latest(self.route_id)

        generate.assert_not_called()
        render_pages.assert_not_called()
        self.assertEqual(result["version_token"], "locked-pdf")
        self.assertEqual(len(result["page_urls"]), 1)

    def test_resolve_file_regenerates_an_unreadable_published_page(self) -> None:
        documents = SopDocumentService(self.store)
        output_dir = documents.root / f"route_{self.route_id}"
        preview_dir = output_dir / CURRENT_PREVIEW_DIR_NAME
        fresh_dir = output_dir / f"{VERSIONED_PREVIEW_DIR_PREFIX}fresh"
        output_dir.mkdir(parents=True)
        preview_dir.mkdir(parents=True)
        fresh_dir.mkdir(parents=True)
        docx_path = output_dir / "current.docx"
        old_pdf = preview_dir / "old.pdf"
        old_page = preview_dir / "page-001.png"
        fresh_pdf = fresh_dir / "fresh.pdf"
        fresh_page = fresh_dir / "page-001.png"
        for path in (docx_path, old_pdf, old_page, fresh_pdf, fresh_page):
            path.write_bytes(path.name.encode("ascii"))
        manifest = {
            "route_id": self.route_id,
            "route_version": 1,
            "product_code": "CHAT-TEST",
            "generated_at": "2026-08-21T00:00:00+00:00",
            "version_token": "old-preview",
            "route_fingerprint": documents._route_fingerprint(self.route_id),
            "template_id": MULTI_PAGE_TEMPLATE_ID,
            "layout_mode": "portrait_flow_then_repeated_landscape_work_instructions",
            "docx_path": str(docx_path),
            "pdf_path": str(old_pdf),
            "page_paths": [str(old_page)],
            "page_count": 1,
            "expected_page_count": 1,
            "validation_path": "",
            "media_count": 0,
            "status": "draft_document_generated",
            "preview_source": "generated_docx",
        }
        manifest_path = output_dir / "document_manifest.json"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        fresh_manifest = manifest | {
            "version_token": "fresh-preview",
            "pdf_path": str(fresh_pdf),
            "page_paths": [str(fresh_page)],
        }

        def regenerate(_: int) -> dict[str, Any]:
            manifest_path.write_text(json.dumps(fresh_manifest), encoding="utf-8")
            return documents.public_manifest(fresh_manifest)

        original_readable = documents._is_readable_file
        with (
            patch.object(
                documents,
                "_is_readable_file",
                side_effect=lambda path: False if Path(path) == old_page else original_readable(Path(path)),
            ),
            patch.object(documents, "generate", side_effect=regenerate) as generate,
        ):
            path, mime_type, filename = documents.resolve_file(self.route_id, "page", page_no=1)

        generate.assert_called_once_with(self.route_id)
        self.assertEqual(path, fresh_page)
        self.assertEqual(mime_type, "image/png")
        self.assertEqual(filename, "page-1.png")

    def test_pdf_page_renderer_supports_chinese_paths(self) -> None:
        import pymupdf

        fixture_dir = Path(self.temp.name) / "中文预览测试"
        output_dir = fixture_dir / "分页图片"
        fixture_dir.mkdir(parents=True)
        pdf_path = fixture_dir / "作业指导书.pdf"
        pdf = pymupdf.open()
        pdf.new_page(width=595, height=842).insert_text((72, 72), "preview page")
        pdf.save(pdf_path)
        pdf.close()

        renderer_script = Path(__file__).resolve().parents[1] / "scripts" / "render_pdf_pages.py"
        result = subprocess.run(
            [
                sys.executable,
                str(renderer_script),
                "--input",
                str(pdf_path),
                "--output-directory",
                str(output_dir),
                "--dpi",
                "110",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )

        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        payload = json.loads(result.stdout)
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(len(payload["page_paths"]), 1)
        self.assertTrue((output_dir / "page-001.png").is_file())


if __name__ == "__main__":
    unittest.main()
